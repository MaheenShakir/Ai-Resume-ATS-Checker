"""
AI Resume ATS Checker
---------------------
Upload a resume (PDF, DOCX or TXT) and get:
  * an ATS score out of 100 with a category breakdown
  * prioritised, concrete improvements
  * missing keywords, strengths and example rewrites

UI: Streamlit  |  AI: Google Gemini Flash (google-genai SDK)
"""

from __future__ import annotations

import io
import json
import os
import re
from typing import Any

import streamlit as st

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #

APP_TITLE = "AI Resume ATS Checker"

# "gemini-flash-latest" is Google's alias that always points at the newest
# Flash model, so the app keeps working when older versions are retired.
DEFAULT_MODEL = "gemini-flash-latest"
FALLBACK_MODELS = ["gemini-flash-latest", "gemini-2.5-flash"]

MAX_FILE_MB = 5
MIN_TEXT_CHARS = 100          # below this the file is probably a scanned image
MAX_RESUME_CHARS = 30_000     # keeps the prompt comfortably inside the context window
MAX_JD_CHARS = 10_000

# Category name -> weight. Weights sum to 1.0.
CATEGORY_WEIGHTS = {
    "Keywords & Relevance": 0.30,
    "Formatting & Parseability": 0.20,
    "Content Quality": 0.20,
    "Structure & Sections": 0.15,
    "Impact & Metrics": 0.15,
}

SYSTEM_PROMPT = """You are an expert technical recruiter and ATS (Applicant Tracking System) \
specialist. You review resumes strictly and honestly. You never invent facts that are not in \
the resume. You respond with a single valid JSON object and nothing else."""


# --------------------------------------------------------------------------- #
# File parsing
# --------------------------------------------------------------------------- #

class ResumeError(Exception):
    """A problem with the uploaded file that we can explain to the user."""


def extract_text_from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise ResumeError("This PDF is password-protected. Please remove the password and retry.") from exc
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n".join(pages)
    except ResumeError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ResumeError("Could not read this PDF. It may be corrupted. Try re-exporting it.") from exc


def extract_text_from_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ResumeError("Could not read this DOCX file. It may be corrupted or not a real .docx.") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    """Return clean text from an uploaded resume, or raise ResumeError."""
    if not data:
        raise ResumeError("The uploaded file is empty.")
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        raise ResumeError(f"File is larger than {MAX_FILE_MB} MB. Please upload a smaller file.")

    name = (filename or "").lower()
    if name.endswith(".pdf"):
        text = extract_text_from_pdf(data)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(data)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ResumeError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    if len(text) < MIN_TEXT_CHARS:
        raise ResumeError(
            "Almost no text could be extracted. If your resume is a scanned image or a "
            "design-heavy PDF, an ATS cannot read it either - export a text-based PDF or DOCX."
        )
    return text


# --------------------------------------------------------------------------- #
# Local (no-AI) sanity checks
# --------------------------------------------------------------------------- #

SECTION_PATTERNS = {
    "Summary / Objective": r"\b(summary|objective|profile|about me)\b",
    "Experience": r"\b(experience|employment|work history|internship)s?\b",
    "Education": r"\b(education|academic|qualification)s?\b",
    "Skills": r"\b(skills|technologies|technical skills|competencies)\b",
    "Projects": r"\bprojects?\b",
    "Certifications": r"\b(certifications?|certificates?|licenses?)\b",
}


def quick_checks(text: str) -> dict[str, Any]:
    lowered = text.lower()
    return {
        "word_count": len(text.split()),
        "has_email": bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)),
        "has_phone": bool(re.search(r"(\+?\d[\d\s().-]{8,}\d)", text)),
        "has_linkedin": "linkedin.com" in lowered,
        "has_github": "github.com" in lowered,
        "sections_found": [name for name, pat in SECTION_PATTERNS.items() if re.search(pat, lowered)],
        "has_numbers": len(re.findall(r"\d+%|\$\s?\d|\d+\+|\b\d{2,}\b", text)) >= 3,
    }


# --------------------------------------------------------------------------- #
# Prompt + response handling
# --------------------------------------------------------------------------- #

def build_prompt(resume_text: str, job_description: str = "") -> str:
    resume_text = resume_text[:MAX_RESUME_CHARS]
    job_description = (job_description or "").strip()[:MAX_JD_CHARS]

    if job_description:
        jd_block = (
            "A target job description is provided. Judge 'Keywords & Relevance' against it, "
            "and fill 'missing_keywords' with important terms from the job description that "
            "are absent from the resume.\n\n"
            f"=== JOB DESCRIPTION ===\n{job_description}\n=== END JOB DESCRIPTION ==="
        )
    else:
        jd_block = (
            "No job description is provided. Infer the most likely target role from the resume "
            "and judge general ATS-friendliness. Fill 'missing_keywords' with commonly expected "
            "keywords for that role that are absent."
        )

    return f"""Analyse the resume below as an ATS would, then return ONE JSON object with exactly \
this structure:

{{
  "target_role": "<role you inferred or the role in the job description>",
  "category_scores": {{
    "Keywords & Relevance": <integer 0-100>,
    "Formatting & Parseability": <integer 0-100>,
    "Content Quality": <integer 0-100>,
    "Structure & Sections": <integer 0-100>,
    "Impact & Metrics": <integer 0-100>
  }},
  "summary": "<2-3 sentence honest overall assessment>",
  "strengths": ["<specific strength>", "..."],
  "improvements": [
    {{"priority": "High" | "Medium" | "Low",
      "section": "<resume section>",
      "issue": "<what is wrong>",
      "fix": "<exactly what to change>"}}
  ],
  "missing_keywords": ["<keyword>", "..."],
  "rewrite_examples": [
    {{"before": "<weak line copied from the resume>", "after": "<stronger version, no invented facts>"}}
  ]
}}

Scoring guidance: be strict and realistic (most resumes score 55-80). Reward quantified \
achievements, action verbs, relevant keywords, clear standard headings, consistent dates and \
a single-column text-based layout. Penalise tables/columns/graphics, missing contact info, \
vague duties, typos, and very long or very short resumes.

Give 5-10 improvements ordered by priority, 3-6 strengths, up to 15 missing keywords and 2-4 \
rewrite examples. Rewrites must only reuse facts already present; use placeholders like [X%] \
where a number is needed.

{jd_block}

=== RESUME ===
{resume_text}
=== END RESUME ===

Return only the JSON object. No markdown fences, no commentary."""


def extract_json(raw: str) -> dict[str, Any]:
    """Pull a JSON object out of a model reply, tolerating code fences / chatter."""
    if not raw or not raw.strip():
        raise ValueError("Empty response from the model.")
    cleaned = raw.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    try:
        obj = json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start == -1 or end <= start:
            raise ValueError("The model did not return JSON.") from None
        try:
            obj = json.loads(cleaned[start : end + 1])
        except json.JSONDecodeError as exc:
            raise ValueError("The model returned malformed JSON.") from exc
    if not isinstance(obj, dict):
        raise ValueError("The model returned an unexpected JSON shape.")
    return obj


def _to_score(value: Any, default: int = 0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError, OverflowError):
        return default


def _to_str_list(value: Any, limit: int = 30) -> list[str]:
    if isinstance(value, str):
        value = [value]
    if not isinstance(value, list):
        return []
    out = [str(v).strip() for v in value if v is not None and str(v).strip()]
    return out[:limit]


def normalize_result(data: dict[str, Any]) -> dict[str, Any]:
    """Validate and clean the model output so the UI can never crash on it."""
    raw_scores = data.get("category_scores")
    raw_scores = raw_scores if isinstance(raw_scores, dict) else {}
    lowered = {str(k).strip().lower(): v for k, v in raw_scores.items()}
    categories = {name: _to_score(lowered.get(name.lower())) for name in CATEGORY_WEIGHTS}

    if not any(categories.values()):
        raise ValueError("The model did not return any category scores.")

    overall = _to_score(sum(categories[n] * w for n, w in CATEGORY_WEIGHTS.items()))

    improvements = []
    order = {"High": 0, "Medium": 1, "Low": 2}
    raw_improvements = data.get("improvements")
    for item in raw_improvements if isinstance(raw_improvements, list) else []:
        if isinstance(item, dict):
            priority = str(item.get("priority", "Medium")).strip().capitalize()
            improvements.append(
                {
                    "priority": priority if priority in order else "Medium",
                    "section": str(item.get("section") or "General").strip(),
                    "issue": str(item.get("issue") or "").strip(),
                    "fix": str(item.get("fix") or "").strip(),
                }
            )
        elif isinstance(item, str) and item.strip():
            improvements.append({"priority": "Medium", "section": "General", "issue": item.strip(), "fix": ""})
    improvements = [i for i in improvements if i["issue"] or i["fix"]]
    improvements.sort(key=lambda i: order[i["priority"]])

    rewrites = []
    raw_rewrites = data.get("rewrite_examples")
    for item in raw_rewrites if isinstance(raw_rewrites, list) else []:
        if isinstance(item, dict) and (item.get("before") or item.get("after")):
            rewrites.append({"before": str(item.get("before") or "").strip(), "after": str(item.get("after") or "").strip()})

    return {
        "overall": overall,
        "categories": categories,
        "target_role": str(data.get("target_role") or "Not detected").strip(),
        "summary": str(data.get("summary") or "").strip(),
        "strengths": _to_str_list(data.get("strengths")),
        "improvements": improvements,
        "missing_keywords": _to_str_list(data.get("missing_keywords")),
        "rewrites": rewrites,
    }


def score_label(score: int) -> tuple[str, str]:
    """Return (label, emoji) for a score."""
    if score >= 80:
        return "Excellent", "🟢"
    if score >= 65:
        return "Good - room to improve", "🟡"
    if score >= 50:
        return "Needs work", "🟠"
    return "Poor - major fixes needed", "🔴"


# --------------------------------------------------------------------------- #
# Gemini call
# --------------------------------------------------------------------------- #

def friendly_api_error(exc: Exception) -> str:
    msg = str(exc)
    low = msg.lower()
    if "api key" in low or "api_key" in low or "401" in low or "403" in low or "permission" in low:
        return "Gemini rejected the API key. Check that it is correct and has access to the Gemini API."
    if "429" in low or "quota" in low or "rate limit" in low or "resource_exhausted" in low:
        return "Gemini rate limit or quota reached. Wait a minute and try again."
    if "503" in low or "unavailable" in low or "overloaded" in low:
        return "Gemini is temporarily overloaded. Please try again in a moment."
    if "timeout" in low or "timed out" in low or "connection" in low:
        return "Could not reach Gemini (network problem). Check your connection and try again."
    return f"Gemini request failed: {msg[:300]}"


def _is_model_not_found(exc: Exception) -> bool:
    low = str(exc).lower()
    return "404" in low or "not_found" in low or "not found" in low or "is not supported" in low


def call_gemini(api_key: str, model: str, prompt: str) -> str:
    """Send the prompt to Gemini and return the raw text reply."""
    try:
        from google import genai
        from google.genai import types
    except ImportError as exc:
        raise RuntimeError("The 'google-genai' package is not installed. Run: pip install -r requirements.txt") from exc

    client = genai.Client(api_key=api_key)
    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_PROMPT,
        temperature=0.2,
        response_mime_type="application/json",
    )

    candidates = []
    for m in [model, *FALLBACK_MODELS]:
        if m and m not in candidates:
            candidates.append(m)

    last_error: Exception | None = None
    for candidate in candidates:
        try:
            response = client.models.generate_content(model=candidate, contents=prompt, config=config)
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            if _is_model_not_found(exc):
                continue  # try the next model name
            raise
        text = getattr(response, "text", None)
        if text:
            return text
        last_error = RuntimeError("Gemini returned an empty response (it may have been blocked).")
    raise last_error or RuntimeError("No Gemini model could be used.")


def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str = "") -> dict[str, Any]:
    """Run the analysis. Retries once if the model returns unusable JSON."""
    prompt = build_prompt(resume_text, job_description)
    last_error: Exception | None = None
    for _ in range(2):
        raw = call_gemini(api_key, model, prompt)
        try:
            return normalize_result(extract_json(raw))
        except ValueError as exc:
            last_error = exc
    raise ValueError(f"Could not understand the AI response ({last_error}). Please try again.")


# --------------------------------------------------------------------------- #
# Report export
# --------------------------------------------------------------------------- #

def build_report(result: dict[str, Any], filename: str) -> str:
    label, _ = score_label(result["overall"])
    lines = [
        "# ATS Resume Report",
        f"File: {filename}",
        f"Target role: {result['target_role']}",
        f"## Overall ATS score: {result['overall']}/100 ({label})",
        "",
        result["summary"],
        "",
        "## Category scores",
    ]
    lines += [f"- {name}: {score}/100" for name, score in result["categories"].items()]
    lines += ["", "## Improvements"]
    for i, imp in enumerate(result["improvements"], 1):
        lines.append(f"{i}. [{imp['priority']}] {imp['section']}: {imp['issue']}")
        if imp["fix"]:
            lines.append(f"   Fix: {imp['fix']}")
    lines += ["", "## Strengths"] + [f"- {s}" for s in result["strengths"]]
    lines += ["", "## Missing keywords", ", ".join(result["missing_keywords"]) or "None"]
    if result["rewrites"]:
        lines += ["", "## Rewrite examples"]
        for r in result["rewrites"]:
            lines += [f"- Before: {r['before']}", f"  After:  {r['after']}"]
    return "\n".join(lines) + "\n"


# --------------------------------------------------------------------------- #
# Streamlit UI
# --------------------------------------------------------------------------- #

def get_default_api_key() -> str:
    """Look for a key in Streamlit secrets, then environment variables."""
    try:
        for name in ("GEMINI_API_KEY", "GOOGLE_API_KEY"):
            value = st.secrets.get(name)
            if value:
                return str(value).strip()
    except Exception:  # noqa: BLE001 - no secrets file is perfectly normal locally
        pass
    return (os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY") or "").strip()


PRIORITY_ICON = {"High": "🔴", "Medium": "🟠", "Low": "🟡"}


def render_results(result: dict[str, Any], checks: dict[str, Any], filename: str) -> None:
    label, emoji = score_label(result["overall"])

    st.divider()
    left, right = st.columns([1, 2])
    with left:
        st.metric("ATS Score", f"{result['overall']} / 100")
        st.markdown(f"**{emoji} {label}**")
        st.caption(f"Target role: {result['target_role']}")
    with right:
        st.progress(result["overall"] / 100)
        if result["summary"]:
            st.write(result["summary"])

    st.subheader("Score breakdown")
    cols = st.columns(len(result["categories"]))
    for col, (name, score) in zip(cols, result["categories"].items()):
        with col:
            st.metric(name, f"{score}")
            st.progress(score / 100)

    tab_fix, tab_kw, tab_good, tab_rewrite, tab_checks = st.tabs(
        ["🛠 Improvements", "🔑 Missing keywords", "✅ Strengths", "✍️ Rewrite examples", "📋 Quick checks"]
    )

    with tab_fix:
        if not result["improvements"]:
            st.info("No specific improvements were returned. Try running the analysis again.")
        for imp in result["improvements"]:
            st.markdown(f"**{PRIORITY_ICON[imp['priority']]} {imp['priority']} - {imp['section']}**")
            if imp["issue"]:
                st.write(f"Issue: {imp['issue']}")
            if imp["fix"]:
                st.write(f"Fix: {imp['fix']}")
            st.write("")

    with tab_kw:
        if result["missing_keywords"]:
            st.write("Consider adding these (only where they are truthful for you):")
            st.write("  ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.success("No important keywords appear to be missing.")

    with tab_good:
        if result["strengths"]:
            for s in result["strengths"]:
                st.markdown(f"- {s}")
        else:
            st.info("No strengths were returned.")

    with tab_rewrite:
        if result["rewrites"]:
            for r in result["rewrites"]:
                st.markdown(f"**Before:** {r['before']}")
                st.markdown(f"**After:** {r['after']}")
                st.write("")
        else:
            st.info("No rewrite examples were returned.")

    with tab_checks:
        def tick(ok: bool) -> str:
            return "✅" if ok else "❌"

        st.write(f"{tick(checks['has_email'])} Email address found")
        st.write(f"{tick(checks['has_phone'])} Phone number found")
        st.write(f"{tick(checks['has_linkedin'])} LinkedIn link found")
        st.write(f"{tick(checks['has_github'])} GitHub link found (optional for non-tech roles)")
        st.write(f"{tick(checks['has_numbers'])} Quantified results (numbers, %, $) found")
        st.write(f"Word count: **{checks['word_count']}** (ideal: roughly 300-800)")
        st.write("Sections detected: " + (", ".join(checks["sections_found"]) or "none"))

    st.download_button(
        "⬇️ Download report",
        data=build_report(result, filename),
        file_name="ats_report.md",
        mime="text/markdown",
    )


def main() -> None:
    st.set_page_config(page_title=APP_TITLE, page_icon="📄", layout="wide")
    st.title("📄 " + APP_TITLE)
    st.caption("Upload your resume to get an ATS score and concrete ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        default_key = get_default_api_key()
        typed_key = st.text_input(
            "Gemini API key",
            type="password",
            help="Get a free key at https://aistudio.google.com/apikey",
            placeholder="Loaded from secrets" if default_key else "Paste your key here",
        )
        model = st.text_input("Gemini model", value=DEFAULT_MODEL, help="Any Gemini Flash model name.")
        st.caption("Your resume is sent to Google's Gemini API for analysis and is not stored by this app.")

    api_key = (typed_key or "").strip() or default_key

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Target job description (optional, recommended)",
        height=150,
        placeholder="Paste the job posting here to get a job-specific score and keyword gaps...",
    )

    if st.button("Analyze resume", type="primary"):
        if uploaded is None:
            st.warning("Please upload a resume first.")
        elif not api_key:
            st.warning("Please enter your Gemini API key in the sidebar.")
        else:
            try:
                resume_text = extract_text(uploaded.name, uploaded.getvalue())
                with st.spinner("Analyzing your resume with Gemini..."):
                    result = analyze_resume(api_key, (model or DEFAULT_MODEL).strip(), resume_text, job_description)
                st.session_state["analysis"] = {
                    "result": result,
                    "checks": quick_checks(resume_text),
                    "filename": uploaded.name,
                }
            except ResumeError as exc:
                st.session_state.pop("analysis", None)
                st.error(str(exc))
            except ValueError as exc:
                st.session_state.pop("analysis", None)
                st.error(str(exc))
            except Exception as exc:  # noqa: BLE001
                st.session_state.pop("analysis", None)
                st.error(friendly_api_error(exc))

    saved = st.session_state.get("analysis")
    if saved:
        render_results(saved["result"], saved["checks"], saved["filename"])


if __name__ == "__main__":
    main()
